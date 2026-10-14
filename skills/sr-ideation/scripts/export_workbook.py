#!/usr/bin/env python3
"""Export agent-authored analysis.json, preserving uncertain scores as text."""
import argparse
import json
import re
from pathlib import Path


SHEETS = [
    ("Prioritas", "ideas", ["id", "topic", "gap", "pico_draft", "unique_reports",
     "independent_studies_confirmed", "independent_studies_potential", "studies_1yr",
     "coverage", "novelty_score", "interest_score", "feasibility_score", "total",
     "confidence", "evidence_ids", "rationale", "next_steps"]),
    ("Studi_Original", "original_studies", ["record_id", "study_id", "title", "year", "doi", "pmid", "source", "eligibility"]),
    ("Existing_SR", "existing_reviews", ["record_id", "title", "year", "last_search_date", "pico", "gap", "source"]),
    ("PICO_Draft", "pico_drafts", ["idea_id", "population", "intervention", "comparator", "outcomes", "design", "status"]),
    ("Audit", "audit", ["item", "detail"]),
    ("Search_Log", "search_log", ["date", "source", "query", "filters", "hits", "retrieved", "limitations"]),
]


def write_value(cell, value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        # Write all strings as strings, including formula-like titles/cells.
        # Excel's cell limit is explicit: full data stays in analysis.json.
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        cell.value = value[:32700] + " [truncated; see analysis.json]" if len(value) > 32700 else value
        cell.data_type = "s"
    else:
        cell.value = value


def export(data, output):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise RuntimeError("openpyxl unavailable: deliver JSON/CSV or use the host spreadsheet facility.") from error
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, key, default_columns in SHEETS:
        rows = data.get(key, [])
        if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
            raise ValueError(f"{key} must be a list of objects")
        if key == "audit" and data.get("phase"):
            rows = [{"item": "phase_decision", "detail": data["phase"]}, *rows]
        columns = default_columns + list(dict.fromkeys(k for row in rows for k in row if k not in default_columns))
        sheet = workbook.create_sheet(title)
        for col, key_name in enumerate(columns, 1):
            header = sheet.cell(1, col, key_name)
            header.fill = PatternFill("solid", fgColor="005447")
            header.font = Font(color="FFFFFF", bold=True)
            sheet.column_dimensions[get_column_letter(col)].width = 26 if key_name not in ("title", "topic", "rationale", "detail") else 48
        for row_index, row in enumerate(rows, 2):
            for col, field in enumerate(columns, 1):
                cell = sheet.cell(row_index, col)
                write_value(cell, row.get(field))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = sheet.dimensions
        else:
            sheet.cell(2, 1, "Belum tersedia; bukan bukti tidak ada studi/review.")
        sheet.sheet_view.showGridLines = False
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("analysis", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    data = json.loads(args.analysis.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        parser.error("analysis.json must contain an object")
    export(data, args.output)
    print(str(args.output.resolve()))


if __name__ == "__main__":
    main()
